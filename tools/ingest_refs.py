#!/usr/bin/env python3
"""대화창에 첨부된 참고 자료를 refs/decks/ 로 받아들이고 텍스트를 추출한다.

사용자가 PPT나 PDF를 첨부하면 업로드 폴더에 들어온다. 이 스크립트는
아직 받아들이지 않은 파일만 골라 복사하고, PPT는 슬라이드별 텍스트를
마크다운으로 뽑아 둔다. 집필 작업자는 그 마크다운을 읽는다.

    python3 tools/ingest_refs.py
"""
import hashlib
import re
import shutil
import subprocess
import sys
from pathlib import Path

UPLOADS = Path("/root/.claude/uploads")
DEST = Path("/home/user/refs/decks")
EXTS = {".pptx", ".ppt", ".pdf", ".docx"}


def digest(p):
    h = hashlib.sha256()
    with p.open("rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()[:16]


def slug(name):
    """업로드 파일명에서 사람이 읽을 이름을 만든다. 한글이 깨져 들어오는 경우가 많다."""
    stem = Path(name).stem
    stem = re.sub(r"^[0-9a-f]{8}-", "", stem)      # 업로드 해시 접두사 제거
    stem = re.sub(r"_{2,}", "-", stem).strip("-_ ")
    stem = re.sub(r"[^0-9A-Za-z가-힣.\- ]", "", stem).strip()
    return stem or "ref"


def to_pptx(src):
    """레거시 .ppt는 python-pptx가 읽지 못하므로 LibreOffice로 변환한다."""
    if src.suffix.lower() != ".ppt":
        return src
    out = src.with_suffix(".pptx")
    if out.exists():
        return out
    subprocess.run(["soffice", "--headless", "--convert-to", "pptx",
                    "--outdir", str(src.parent), str(src)],
                   capture_output=True, timeout=300)
    return out if out.exists() else src


def extract_pptx(src, out_md):
    from pptx import Presentation
    src = to_pptx(src)
    p = Presentation(str(src))
    slides = list(p.slides)
    lines = [f"# {out_md.stem}", "", f"- 원본: `refs/decks/{src.name}`",
             f"- 슬라이드 {len(slides)}장", ""]
    for i, s in enumerate(slides, 1):
        txt = [sh.text_frame.text.strip() for sh in s.shapes
               if sh.has_text_frame and sh.text_frame.text.strip()]
        if not txt:
            continue
        lines.append(f"## {i}. {txt[0].splitlines()[0]}")
        for t in txt[1:]:
            lines.append(t.replace("\n", "  \n"))
        lines.append("")
    out_md.write_text("\n".join(lines), encoding="utf-8")
    return len(slides)


def extract_docx(src, out_md):
    from docx import Document
    d = Document(str(src))
    lines = [f"# {out_md.stem}", "", f"- 원본: `refs/decks/{src.name}`", ""]
    for para in d.paragraphs:
        t = para.text.strip()
        if not t:
            continue
        style = (para.style.name or "").lower()
        if "heading" in style:
            lvl = "".join(c for c in style if c.isdigit()) or "2"
            lines.append("#" * min(int(lvl) + 1, 6) + " " + t)
        else:
            lines.append(t)
    for tb in d.tables:
        for row in tb.rows:
            lines.append(" | ".join(c.text.strip().replace("\n", " ") for c in row.cells))
    out_md.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return len(d.paragraphs)


def extract_pdf(src, out_md):
    r = subprocess.run(["pdftotext", "-layout", str(src), "-"],
                       capture_output=True, text=True)
    if r.returncode != 0:
        return 0
    out_md.write_text(f"# {out_md.stem}\n\n- 원본: `refs/decks/{src.name}`\n\n```\n"
                      + r.stdout[:400000] + "\n```\n", encoding="utf-8")
    return r.stdout.count("\f") + 1


def main():
    DEST.mkdir(parents=True, exist_ok=True)
    known = {digest(p) for p in DEST.iterdir() if p.suffix.lower() in EXTS}
    found = sorted(q for d in UPLOADS.iterdir() if d.is_dir()
                   for q in d.iterdir() if q.suffix.lower() in EXTS)
    new = 0
    for src in found:
        if digest(src) in known:
            continue
        name = slug(src.name)
        dst = DEST / f"{name}{src.suffix.lower()}"
        n = 1
        while dst.exists():
            n += 1
            dst = DEST / f"{name}-{n}{src.suffix.lower()}"
        shutil.copy2(src, dst)
        md = dst.with_suffix(".md")
        try:
            if dst.suffix in (".pptx", ".ppt"):
                pages = extract_pptx(dst, md)
            elif dst.suffix == ".docx":
                pages = extract_docx(dst, md)
            else:
                pages = extract_pdf(dst, md)
        except Exception as e:
            pages = 0
            print(f"  텍스트 추출 실패 {dst.name}: {e}", file=sys.stderr)
        print(f"새 자료: {dst.name}  {pages}쪽  → {md.name}")
        new += 1
    if not new:
        print("새로 들어온 자료가 없습니다.")
    else:
        print(f"\n{new}건을 받아들였습니다. refs/decks/INDEX.md에 대응 수업을 적어 두세요.")


if __name__ == "__main__":
    main()
