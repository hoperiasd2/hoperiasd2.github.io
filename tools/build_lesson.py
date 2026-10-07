#!/usr/bin/env python3
"""수업 원본(lesson.md / protocol.md)에서 index.html과 slides.pptx를 생성한다.

하나의 원본이 두 산출물의 공통 명세다. 수치·용어가 갈라지지 않도록
본문을 두 곳에 따로 쓰지 않는다.

사용:
    python3 tools/build_lesson.py M02-W09-L03
    python3 tools/build_lesson.py --all          원본이 있는 모든 수업
    python3 tools/build_lesson.py M02-W09-L03 --html-only
"""
import argparse
import csv
import html
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent

# --------------------------------------------------------------------------- 파싱

@dataclass
class Sec:
    title: str = ""
    question: str = ""
    answer: str = ""
    blocks: list = field(default_factory=list)


def _inline(t):
    """**굵게**, *기울임*, `코드`, ^{위}, _{아래} → HTML"""
    t = html.escape(t, quote=False)
    t = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", t)
    t = re.sub(r"(?<!\w)\*([^*]+?)\*(?!\w)", r"<em>\1</em>", t)
    t = re.sub(r"`(.+?)`", r"<code>\1</code>", t)
    t = re.sub(r"\^\{(.+?)\}", r"<sup>\1</sup>", t)
    t = re.sub(r"_\{(.+?)\}", r"<sub>\1</sub>", t)
    return t


def plain(t):
    t = re.sub(r"\*\*(.+?)\*\*", r"\1", t)
    t = re.sub(r"(?<!\w)\*([^*]+?)\*(?!\w)", r"\1", t)
    t = re.sub(r"`(.+?)`", r"\1", t)
    t = re.sub(r"\^\{(.+?)\}", r"\1", t)
    t = re.sub(r"_\{(.+?)\}", r"\1", t)
    return t


def parse_blocks(lines):
    """문단, 불릿(- , 2칸 들여쓰기로 하위), 표(| |), 조건상자(> ), 귀결상자(=> )"""
    blocks, buf, mode = [], [], None

    def flush():
        nonlocal buf, mode
        if buf:
            blocks.append((mode, buf))
        buf, mode = [], None

    for ln in lines:
        s = ln.rstrip()
        if not s.strip():
            flush()
            continue
        if s.startswith("=> "):
            flush()
            blocks.append(("result", [s[3:].strip()]))
        elif s.startswith("> "):
            flush()
            blocks.append(("limit", [s[2:].strip()]))
        elif s.lstrip().startswith("- "):
            lvl = 1 if s.startswith("  - ") else 0
            if mode != "ul":
                flush()
                mode = "ul"
            buf.append((lvl, s.lstrip()[2:].strip()))
        elif s.startswith("@fig "):
            flush()
            parts = s[5:].split(None, 1)
            blocks.append(("fig", [parts[0], parts[1] if len(parts) > 1 else ""]))
        elif s.startswith("|"):
            if mode != "table":
                flush()
                mode = "table"
            cells = [c.strip() for c in s.strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                buf.append(cells)
        else:
            if mode != "p":
                flush()
                mode = "p"
            buf.append(s.strip())
    flush()
    return blocks


def parse_source(path):
    text = path.read_text(encoding="utf-8")
    meta, body = {}, text
    if text.startswith("---"):
        _, fm, body = text.split("---", 2)
        for ln in fm.strip().split("\n"):
            if ":" in ln:
                k, v = ln.split(":", 1)
                meta[k.strip()] = v.strip()

    named = {}          # @obj, @prereq, @safety, @materials, @record, @ref 등
    secs, quiz = [], []
    cur_key, cur_lines, cur_sec = None, [], None

    def close():
        nonlocal cur_key, cur_lines, cur_sec
        if cur_sec is not None:
            cur_sec.blocks = parse_blocks(cur_lines)
            secs.append(cur_sec)
            cur_sec = None
        elif cur_key:
            named[cur_key] = cur_lines[:]
        cur_key, cur_lines = None, []

    for ln in body.split("\n"):
        m = re.match(r"^@(\w+)\s*(.*)$", ln)
        if m and m.group(1) == "fig":
            m = None          # @fig는 절 안의 블록이므로 상위 태그로 보지 않는다
        if m:
            close()
            tag, rest = m.group(1), m.group(2).strip()
            if tag == "sec":
                cur_sec = Sec(title=rest)
            else:
                cur_key = tag
            continue
        if cur_sec is not None:
            if ln.startswith("?") and not cur_sec.question:
                cur_sec.question = ln[1:].strip().lstrip(":").strip()
                continue
            if ln.startswith("!") and not cur_sec.answer:
                cur_sec.answer = ln[1:].strip().lstrip(":").strip()
                continue
        cur_lines.append(ln)
    close()

    # 확인 문항
    q = None
    for ln in named.get("quiz", []):
        if ln.startswith("Q:"):
            q = ln[2:].strip()
        elif ln.startswith("A:") and q:
            quiz.append((q, ln[2:].strip()))
            q = None

    lists = {k: [l.strip()[2:].strip() for l in v if l.strip().startswith("- ")]
             for k, v in named.items()}
    return meta, named, lists, secs, quiz


# --------------------------------------------------------------------------- HTML

def esc(s):
    return html.escape(str(s), quote=True)


FIGN = {"n": 0}
FIG_SRC = {}
FIG_CREDIT = {}
FIG_EXT = (".svg", ".png", ".jpg", ".jpeg", ".webp", ".gif")


def find_figure(lesson_dir, fid):
    """assets/ 에서 그림 파일을 찾는다. SVG와 래스터 이미지를 모두 지원한다."""
    for ext in FIG_EXT:
        f = Path(lesson_dir) / "assets" / f"{fid}{ext}"
        if f.exists():
            return f
    return None


def render_blocks(blocks):
    out = []
    for kind, buf in blocks:
        if kind == "fig":
            FIGN["n"] += 1
            fid, cap = buf[0], buf[1]
            src = FIG_SRC.get(fid, f"assets/{fid}.svg")
            credit = FIG_CREDIT.get(fid, "")
            cr = f' <span class="credit">{_inline(credit)}</span>' if credit else ""
            out.append(
                f'<figure class="fig" id="{esc(fid)}">'
                f'<img src="{esc(src)}" alt="{esc(plain(cap))}" loading="lazy">'
                f'<figcaption><b>그림 {FIGN["n"]}</b> {_inline(cap)}{cr}</figcaption></figure>')
            continue
        if kind == "p":
            out.append("<p>" + _inline(" ".join(buf)) + "</p>")
        elif kind == "ul":
            depth, parts = 0, ["<ul>"]
            for lvl, item in buf:
                if lvl > depth:
                    parts.append("<ul>")
                    depth = lvl
                while lvl < depth:
                    parts.append("</ul>")
                    depth -= 1
                parts.append("<li>" + _inline(item) + "</li>")
            parts.append("</ul>" * (depth + 1))
            out.append("".join(parts))
        elif kind == "table":
            head, *rest = buf
            th = "".join(f"<th>{_inline(c)}</th>" for c in head)
            tr = "".join("<tr>" + "".join(f"<td>{_inline(c)}</td>" for c in r) + "</tr>"
                         for r in rest)
            out.append(f'<div class="tw"><table><thead><tr>{th}</tr></thead><tbody>{tr}</tbody></table></div>')
        elif kind == "limit":
            out.append('<div class="box box--limit"><b>임상 연계</b> ' + _inline(buf[0]) + "</div>")
        elif kind == "result":
            out.append('<div class="box box--result"><b>핵심 정리</b> ' + _inline(buf[0]) + "</div>")
    return "\n  ".join(out)


SECTION_LABEL = {
    "prereq": "선수지식", "obj": "학습목표", "safety": "안전·윤리 고지",
    "materials": "준비물·시약·장비", "record": "결과 기록", "trouble": "문제해결",
}


def lesson_html(row, meta, named, lists, secs, quiz, prev_r, next_r, lesson_dir=None):
    FIGN["n"] = 0
    FIG_SRC.clear()
    if lesson_dir:
        for sec in secs:
            for kind, buf in sec.blocks:
                if kind == "fig":
                    f = find_figure(lesson_dir, buf[0])
                    if f:
                        FIG_SRC[buf[0]] = f"assets/{f.name}"
    practicum = row["kind"] == "practicum"
    up = "../../../../"

    def crumb():
        parts = [("대문", up + "index.html"), ("전체 수업", up + "curriculum/index.html"),
                 (row["module"], "../../index.html"), (f'{row["week"]}주차', "../index.html"),
                 (row["lesson_id"], None)]
        o = ['<nav class="crumb" aria-label="이동 경로">']
        for i, (lab, href) in enumerate(parts):
            if i:
                o.append('<span aria-hidden="true">›</span>')
            o.append(f'<a href="{href}">{esc(lab)}</a>' if href else f"<span>{esc(lab)}</span>")
        o.append("</nav>")
        return "".join(o)

    body = [crumb(), f"""
  <section class="hero">
    <p class="eyebrow">{row["lesson_id"]} · {esc(row["module_title"])}{" · " + esc(row["group"]) if row["group"] else ""}</p>
    <h1>{esc(meta.get("title", row["title"]))}</h1>
    <p class="en">{esc(meta.get("en", ""))}</p>
  </section>"""]

    if (meta.get("slides") or "").lower() != "none":
        body.append('  <p class="dl"><a class="dlbtn" href="slides.pptx">슬라이드 내려받기 (PPTX)</a></p>')

    # 안전 고지는 실습에서 맨 위에 둔다.
    if practicum and lists.get("safety"):
        items = "".join(f"<li>{_inline(x)}</li>" for x in lists["safety"])
        body.append(f'  <div class="box box--safety"><b>안전·윤리 고지</b><ul>{items}</ul></div>')

    for key in ("obj", "prereq"):
        if lists.get(key):
            items = "".join(f"<li>{_inline(x)}</li>" for x in lists[key])
            body.append(f'  <h2>{SECTION_LABEL[key]}</h2>\n  <ul>{items}</ul>')

    if secs:
        toc = "".join(
            f'<li><a href="#s{i+1}">{_inline(s.title)}</a></li>'
            for i, s in enumerate(secs))
        label = "실습 구성" if practicum else "강의 구성"
        body.append(f'  <h2>{label}</h2>\n  <ol class="qlist">{toc}</ol>')

    if practicum and lists.get("materials"):
        items = "".join(f"<li>{_inline(x)}</li>" for x in lists["materials"])
        body.append(f'  <h2>준비물·시약·장비</h2>\n  <ul>{items}</ul>')

    for i, s in enumerate(secs, 1):
        body.append(f'  <section id="s{i}">')
        body.append(f'  <h2>{_inline(s.title)}</h2>')
        if s.question:
            body.append(f'  <p class="q"><span>질문</span>{_inline(s.question)}</p>')
        if s.answer:
            body.append(f'  <p class="a"><span>답</span>{_inline(s.answer)}</p>')
        body.append("  " + render_blocks(s.blocks))
        body.append("  </section>")

    if named.get("trouble"):
        body.append("  <h2>문제해결</h2>\n  " + render_blocks(parse_blocks(named["trouble"])))
    if named.get("record"):
        body.append("  <h2>결과 기록</h2>\n  " + render_blocks(parse_blocks(named["record"])))

    if quiz:
        items = "".join(
            f'<li><p class="qq">{_inline(q)}</p>'
            f'<details><summary>해설</summary><p>{_inline(a)}</p></details></li>'
            for q, a in quiz)
        body.append(f'  <h2>확인 문항</h2>\n  <ol class="quiz">{items}</ol>')

    if lists.get("ref"):
        items = "".join(f"<li>{_inline(x)}</li>" for x in lists["ref"])
        body.append(f'  <h2>참고문헌</h2>\n  <ol class="ref">{items}</ol>')

    pg = []
    pg.append(f'<a href="../../../{prev_r["path"][len("curriculum/"):]}index.html">← {prev_r["lesson_id"]}</a>'
              if prev_r else "<span>처음 수업</span>")
    pg.append(f'<a href="../../../{next_r["path"][len("curriculum/"):]}index.html">{next_r["lesson_id"]} →</a>'
              if next_r else "<span>마지막 수업</span>")
    body.append(f'  <nav class="pager">{pg[0]}{pg[1]}</nav>')

    inner = "\n".join(body)
    return f"""<!DOCTYPE html>
<html lang="ko">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{esc(meta.get("title", row["title"]))} · {row["lesson_id"]}</title>
<meta name="description" content="{esc(plain(meta.get("title", row["title"])))}">
<link rel="stylesheet" href="{up}assets/site.css">
<link rel="stylesheet" href="{up}assets/lesson.css">
</head>
<body>

<header class="masthead">
  <div class="wrap">
    <a class="brand" href="{up}index.html">의생명과학 교육과정</a>
    <nav>
      <a href="{up}curriculum/index.html">전체 수업</a>
      <a href="{up}curriculum/materials.html">자료 내려받기</a>
      <a href="{up}archive/index.html">아카이브</a>
    </nav>
  </div>
</header>

<main class="wrap wrap--narrow lesson">
{inner}
</main>

<footer class="foot wrap">
  <span>© 2026 Kim Jintae · Hanyang University</span>
  <a href="https://github.com/hoperiasd2/hoperiasd2.github.io">GitHub 저장소</a>
</footer>
</body>
</html>
"""


# --------------------------------------------------------------------------- PPTX

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

CHROME_BIN = next((b for b in (
    "/opt/pw-browsers/chromium-1194/chrome-linux/chrome",
    "/usr/bin/chromium", "/usr/bin/chromium-browser", "/usr/bin/google-chrome",
) if Path(b).exists()), None)


def rasterize(svg_path, out_png, width=2000):
    """Chromium으로 SVG를 PNG로 변환한다. 성공하면 True."""
    import re as _re
    import subprocess
    import tempfile
    if not CHROME_BIN:
        return False
    svg = Path(svg_path).read_text(encoding="utf-8")
    m = _re.search(r'viewBox="\s*[\d.+-]+\s+[\d.+-]+\s+([\d.]+)\s+([\d.]+)', svg)
    if not m:
        return False
    vw, vh = float(m.group(1)), float(m.group(2))
    h = int(round(width * vh / vw))
    with tempfile.TemporaryDirectory() as td:
        page = Path(td) / "p.html"
        page.write_text(
            '<!doctype html><meta charset="utf-8">'
            '<style>html,body{margin:0;background:#fff}img{display:block;width:%dpx}</style>'
            '<img src="file://%s">' % (width, Path(svg_path).resolve()), encoding="utf-8")
        r = subprocess.run(
            [CHROME_BIN, "--headless", "--disable-gpu", "--no-sandbox", "--hide-scrollbars",
             f"--window-size={width},{h}", f"--screenshot={out_png}", f"file://{page}"],
            capture_output=True, timeout=120)
    return Path(out_png).exists() and r.returncode == 0


NAVY = RGBColor(0x00, 0x3C, 0x71)
CHROME = RGBColor(0xF2, 0xB7, 0x05)
INK = RGBColor(0x11, 0x20, 0x2B)
MUTED = RGBColor(0x66, 0x78, 0x8A)
PAPER = RGBColor(0xF5, 0xF8, 0xF9)
QBG = RGBColor(0xE8, 0xF1, 0xFA)
ABG = RGBColor(0xFF, 0xF5, 0xD6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Malgun Gothic"
W, H = Inches(13.333), Inches(7.5)


def set_font(run, size, bold=False, color=INK, italic=False):
    f = run.font
    f.size, f.bold, f.italic, f.name = Pt(size), bold, italic, FONT
    f.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", FONT)


def add_rich(p, text, size, color=INK, bold=False):
    for part in re.split(r"(\*\*.+?\*\*)", text):
        if not part:
            continue
        b = bold
        if part.startswith("**") and part.endswith("**"):
            part, b = part[2:-2], True
        r = p.add_run()
        r.text = plain(part)
        set_font(r, size, bold=b, color=color)


def rect(slide, x, y, w, h, fill, line=None):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
    s.shadow.inherit = False
    return s


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tf = slide.shapes.add_textbox(x, y, w, h).text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.1)
    tf.margin_top = tf.margin_bottom = Inches(0.05)
    return tf


class Deck:
    def __init__(self, row, meta):
        self.row, self.meta = row, meta
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = W, H
        self.blank = self.prs.slide_layouts[6]
        self.n = 0

    def _base(self, title, kicker=""):
        s = self.prs.slides.add_slide(self.blank)
        self.n += 1
        rect(s, 0, 0, W, H, PAPER)
        rect(s, 0, 0, W, Inches(1.05), NAVY)
        rect(s, 0, Inches(1.05), W, Inches(0.06), CHROME)
        tf = textbox(s, Inches(0.5), Inches(0.12), Inches(12.3), Inches(0.85), MSO_ANCHOR.MIDDLE)
        p = tf.paragraphs[0]
        if kicker:
            r = p.add_run()
            r.text = kicker + "   "
            set_font(r, 13, bold=True, color=CHROME)
        add_rich(p, title, 24 if len(plain(title)) < 36 else 20, color=WHITE, bold=True)
        ftf = textbox(s, Inches(0.5), Inches(7.05), Inches(12.3), Inches(0.33), MSO_ANCHOR.MIDDLE)
        fp = ftf.paragraphs[0]
        r = fp.add_run()
        r.text = f'{self.row["lesson_id"]}  ·  {plain(self.meta.get("title", self.row["title"]))}'
        set_font(r, 9, color=MUTED)
        r2 = fp.add_run()
        r2.text = f"      {self.n}"
        set_font(r2, 9, bold=True, color=NAVY)
        return s

    def note(self, slide, text):
        if text:
            slide.notes_slide.notes_text_frame.text = plain(text)

    def title_slide(self):
        s = self.prs.slides.add_slide(self.blank)
        self.n += 1
        rect(s, 0, 0, W, H, NAVY)
        rect(s, 0, Inches(4.9), W, Inches(0.08), CHROME)
        tf = textbox(s, Inches(0.8), Inches(1.0), Inches(11.7), Inches(0.6))
        r = tf.paragraphs[0].add_run()
        r.text = f'{self.row["module"]}  {self.row["module_title"]}'
        set_font(r, 14, bold=True, color=CHROME)
        tf = textbox(s, Inches(0.8), Inches(1.8), Inches(11.7), Inches(2.4), MSO_ANCHOR.BOTTOM)
        add_rich(tf.paragraphs[0], self.meta.get("title", self.row["title"]), 34, color=WHITE, bold=True)
        if self.meta.get("en"):
            p = tf.add_paragraph()
            r = p.add_run()
            r.text = self.meta["en"]
            set_font(r, 15, color=RGBColor(0xB9, 0xD2, 0xE8))
        tf = textbox(s, Inches(0.8), Inches(5.2), Inches(11.7), Inches(1.2))
        r = tf.paragraphs[0].add_run()
        r.text = self.row["lesson_id"]
        set_font(r, 13, bold=True, color=WHITE)
        p = tf.add_paragraph()
        r = p.add_run()
        r.text = f'{self.row["week"]}주 {self.row["session"]}회' + (f'  ·  {self.row["group"]}' if self.row["group"] else "")
        set_font(r, 11, color=RGBColor(0x9F, 0xBE, 0xDA))
        return s

    def bullets(self, title, items, kicker="", note="", numbered=False):
        """items: [(level, text)] 또는 [text]. 분량이 많으면 자동 분할."""
        norm = [(0, x) if isinstance(x, str) else x for x in items]
        pages, cur, load = [], [], 0
        for lvl, t in norm:
            c = len(plain(t))
            if cur and (load + c > 760 or len(cur) >= 9):
                pages.append(cur)
                cur, load = [], 0
            cur.append((lvl, t))
            load += c
        if cur:
            pages.append(cur)
        for i, pg in enumerate(pages):
            ttl = title if len(pages) == 1 else f"{title} ({i+1}/{len(pages)})"
            s = self._base(ttl, kicker)
            tf = textbox(s, Inches(0.7), Inches(1.4), Inches(11.9), Inches(5.5))
            dens = sum(len(plain(t)) for _, t in pg)
            fs = 20 if dens < 240 else (18 if dens < 420 else (16 if dens < 600 else 14))
            for j, (lvl, t) in enumerate(pg):
                p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                p.space_after = Pt(9 if lvl == 0 else 5)
                r = p.add_run()
                r.text = ("•  " if lvl == 0 else "–  ") if not numbered else f"{j+1}.  "
                set_font(r, fs, bold=lvl == 0, color=CHROME if lvl == 0 else MUTED)
                add_rich(p, t, fs if lvl == 0 else fs - 2,
                         color=INK if lvl == 0 else MUTED)
                p.level = lvl
            if i == 0:
                self.note(s, note)

    def qa(self, idx, sec):
        """절의 질문과 답을 한 장으로."""
        s = self._base(plain(sec.title), f"질문 {idx}")
        box = rect(s, Inches(0.7), Inches(1.45), Inches(11.9), Inches(2.0), QBG)
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = Inches(0.3)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        r = p.add_run()
        r.text = "질문   "
        set_font(r, 13, bold=True, color=NAVY)
        add_rich(p, sec.question, 18, color=INK)
        box = rect(s, Inches(0.7), Inches(3.65), Inches(11.9), Inches(2.6), ABG)
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = Inches(0.3)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        r = p.add_run()
        r.text = "답   "
        set_font(r, 13, bold=True, color=RGBColor(0x8A, 0x63, 0x00))
        add_rich(p, sec.answer, 19, color=INK, bold=True)
        self.note(s, f"질문: {plain(sec.question)}\n답: {plain(sec.answer)}")

    def table(self, title, rows, kicker=""):
        head, body = rows[0], rows[1:]
        ncol = len(head)
        for i in range(0, max(1, len(body)), 8):
            chunk = body[i:i + 8]
            n = len(body)
            ttl = title if n <= 8 else f"{title} ({i//8+1}/{(n+7)//8})"
            s = self._base(ttl, kicker)
            data = [head] + chunk
            shp = s.shapes.add_table(len(data), ncol, Inches(0.6), Inches(1.4),
                                     Inches(12.1), Inches(0.4) * len(data)).table
            for ri, row in enumerate(data):
                for ci in range(ncol):
                    cell = shp.cell(ri, ci)
                    cell.text = ""
                    p = cell.text_frame.paragraphs[0]
                    add_rich(p, row[ci] if ci < len(row) else "", 12 if ri else 12,
                             color=WHITE if ri == 0 else INK, bold=ri == 0)
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = NAVY if ri == 0 else (
                        WHITE if ri % 2 else RGBColor(0xEF, 0xF4, 0xF8))
                    cell.margin_left = cell.margin_right = Inches(0.08)
                    cell.margin_top = cell.margin_bottom = Inches(0.03)

    def boxes(self, title, items, kicker=""):
        """items: [(kind, text)] — kind는 limit 또는 result"""
        s = self._base(title, kicker)
        y, avail = Inches(1.45), Inches(5.45)
        weights = [max(len(plain(t)), 40) for _, t in items]
        for (kind, t), wgt in zip(items, weights):
            share = Emu(int(avail * wgt / sum(weights))) - Inches(0.15)
            fs = 16
            lines = max(1, len(plain(t)) // 72 + 1)
            need = Pt(lines * fs * 1.45) + Inches(0.55)
            h = min(share, need)
            box = rect(s, Inches(0.7), y, Inches(11.9), h,
                       QBG if kind == "limit" else ABG)
            tf = box.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.margin_left = tf.margin_right = Inches(0.25)
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            r = p.add_run()
            r.text = ("임상 연계   " if kind == "limit" else "핵심 정리   ")
            set_font(r, 12, bold=True, color=NAVY if kind == "limit" else RGBColor(0x8A, 0x63, 0x00))
            add_rich(p, t, fs, color=INK)
            y += h + Inches(0.15)

    def figure(self, title, svg_path, caption, kicker=""):
        """SVG를 PNG로 변환해 한 장에 싣는다.

        한글 글꼴이 fontconfig에 없으므로 Chromium으로 래스터화한다.
        Chromium을 쓸 수 없으면 cairosvg로 되돌아간다(한글이 깨질 수 있다)."""
        src = Path(svg_path)
        tmp = None
        if src.suffix.lower() == ".svg":
            png = str(src).replace(".svg", ".slide.png")
            tmp = png
            if not rasterize(src, png, width=2200):
                import cairosvg
                cairosvg.svg2png(url=str(src), write_to=png, scale=2.0, background_color="white")
        else:
            png = str(src)
        s = self._base(title, kicker)
        from PIL import Image
        iw, ih = Image.open(png).size
        maxw, maxh = Inches(11.6), Inches(4.75)
        scale = min(maxw / iw, maxh / ih)
        w, h = int(iw * scale), int(ih * scale)
        s.shapes.add_picture(png, int((W - w) / 2), Inches(1.45), w, h)
        if caption:
            tf = textbox(s, Inches(0.7), Inches(6.35), Inches(11.9), Inches(0.6))
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            add_rich(p, caption, 12, color=MUTED)
        self.note(s, plain(caption))
        if tmp:
            Path(tmp).unlink(missing_ok=True)

    def quiz(self, quiz):
        for i, (q, a) in enumerate(quiz, 1):
            s = self._base("확인 문항", f"Q{i}")
            tf = textbox(s, Inches(0.7), Inches(1.5), Inches(11.9), Inches(2.3))
            add_rich(tf.paragraphs[0], q, 19, color=INK, bold=True)
            box = rect(s, Inches(0.7), Inches(4.0), Inches(11.9), Inches(2.85), WHITE,
                       line=RGBColor(0xC2, 0xCD, 0xD6))
            tf = box.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_right = Inches(0.25)
            tf.margin_top = Inches(0.18)
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            r = p.add_run()
            r.text = "해설   "
            set_font(r, 12, bold=True, color=NAVY)
            add_rich(p, a, 15, color=INK)
            self.note(s, plain(a))


def build_deck(row, meta, lists, secs, quiz, out, lesson_dir):
    d = Deck(row, meta)
    d.title_slide()
    if lists.get("obj"):
        d.bullets("학습목표", lists["obj"], kicker="GOALS", numbered=True)
    if lists.get("prereq"):
        d.bullets("선수지식", lists["prereq"], kicker="PREREQ")
    if row["kind"] == "practicum" and lists.get("safety"):
        d.bullets("안전·윤리 고지", lists["safety"], kicker="SAFETY")
    if row["kind"] == "practicum" and lists.get("materials"):
        d.bullets("준비물·시약·장비", lists["materials"], kicker="SETUP")
    if secs:
        d.bullets("강의 구성", [s.title for s in secs], kicker="OUTLINE", numbered=True)
    for i, sec in enumerate(secs, 1):
        if sec.question and sec.answer:
            d.qa(i, sec)
        kick = f"{i:02d}"
        pend, text = [], []

        def flush_text():
            if text:
                d.bullets(plain(sec.title), list(text), kick)
                text.clear()

        for kind, buf in sec.blocks:
            if kind in ("limit", "result"):
                pend.append((kind, buf[0]))
                continue
            if pend:
                flush_text()
                d.boxes(plain(sec.title), pend, kick)
                pend = []
            if kind == "ul":
                text.extend(buf)
            elif kind == "p":
                text.append((0, " ".join(buf)))
            elif kind == "table":
                flush_text()
                d.table(plain(sec.title), buf, kick)
            elif kind == "fig":
                flush_text()
                f = find_figure(lesson_dir, buf[0])
                if f:
                    d.figure(plain(sec.title), f, buf[1], kick)
        flush_text()
        if pend:
            d.boxes(plain(sec.title), pend, kick)
    if quiz:
        d.quiz(quiz)
    if lists.get("ref"):
        d.bullets("참고문헌", lists["ref"], kicker="REFERENCES")
    d.prs.save(out)
    return d.n


# --------------------------------------------------------------------------- 실행

def load_rows():
    with (ROOT / "curriculum.csv").open(encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    for r in rows:
        r["week"], r["session"] = int(r["week"]), int(r["session"])
    return rows


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("lesson_id", nargs="?")
    ap.add_argument("--all", action="store_true")
    ap.add_argument("--html-only", action="store_true")
    a = ap.parse_args()

    rows = load_rows()
    by_id = {r["lesson_id"]: r for r in rows}

    if a.all:
        targets = []
        for r in rows:
            src = ROOT / r["path"] / ("protocol.md" if r["kind"] == "practicum" else "lesson.md")
            if src.exists() and "작성 예정" not in src.read_text(encoding="utf-8")[:4000]:
                targets.append(r["lesson_id"])
    elif a.lesson_id:
        targets = [a.lesson_id]
    else:
        sys.exit("수업 ID 또는 --all이 필요합니다.")

    if not targets:
        print("집필된 원본이 없습니다.")
        return

    for lid in targets:
        row = by_id.get(lid)
        if not row:
            sys.exit(f"{lid}: curriculum.csv에 없는 수업 ID")
        d = ROOT / row["path"]
        src = d / ("protocol.md" if row["kind"] == "practicum" else "lesson.md")
        if not src.exists():
            sys.exit(f"{lid}: 원본 {src.name} 없음")
        meta, named, lists, secs, quiz = parse_source(src)

        i = rows.index(row)
        prev_r = rows[i - 1] if i and rows[i - 1]["module"] == row["module"] else None
        next_r = rows[i + 1] if i + 1 < len(rows) and rows[i + 1]["module"] == row["module"] else None

        (d / "index.html").write_text(
            lesson_html(row, meta, named, lists, secs, quiz, prev_r, next_r, d), encoding="utf-8")

        n = 0
        if not a.html_only and (meta.get("slides") or "").lower() != "none":
            n = build_deck(row, meta, lists, secs, quiz, str(d / "slides.pptx"), d)
        figs = []
        k = 0
        for sec in secs:
            for kind, buf in sec.blocks:
                if kind == "fig":
                    k += 1
                    figs.append({
                        "figure_id": buf[0], "lesson_id": lid,
                        "html_anchor": f"#{buf[0]}", "caption": buf[1],
                        "asset_path": (lambda f: f"assets/{f.name}" if f else "")(find_figure(d, buf[0])),
                        "source_id": "", "rights_status": "자체 제작",
                        "status": "확정" if find_figure(d, buf[0]) else "미제작",
                        "replacement_requirements": "", "review_result": "",
                    })
        (d / "figures.json").write_text(
            __import__("json").dumps({"lesson_id": lid, "figures": figs},
                                     ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        print(f"{lid}  절 {len(secs)}  그림 {len(figs)}  문항 {len(quiz)}  슬라이드 {n}")


if __name__ == "__main__":
    main()
